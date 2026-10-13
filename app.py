import io
import json
import os
import re
import tempfile
import time
import urllib.request
import uuid
from datetime import datetime

import ctranslate2
import streamlit as st
from docx import Document
from fpdf import FPDF
from pypdf import PdfReader
from huggingface_hub import snapshot_download
from transformers import AutoTokenizer

# Optional ChromaDB
try:
    import chromadb
    from chromadb.config import Settings
    HAS_CHROMA = True
except ImportError:
    HAS_CHROMA = False

# ------------------------------------------------------------------ config
MODEL_NAME = "olob0/nllb-200-distilled-600M-ct2-int8_float16"
MAX_CHUNK_CHARS = 300
MAX_UPLOAD_MB = 5
LARGE_THRESHOLD = 100      # more unique chunks than this -> batched mode
BIG_BATCH = 8
DISPLAY_LIMIT = 20000      # max characters shown in the output box
CHROMA_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "chroma_db")

LANGUAGES = {
    "English": "eng_Latn",
    "Hindi": "hin_Deva",
    "Telugu": "tel_Telu",
    "Tamil": "tam_Taml",
    "Kannada": "kan_Knda",
    "Malayalam": "mal_Mlym",
    "Bengali": "ben_Beng",
    "Marathi": "mar_Deva",
    "Urdu": "urd_Arab",
    "French": "fra_Latn",
    "German": "deu_Latn",
    "Spanish": "spa_Latn",
    "Italian": "ita_Latn",
    "Arabic": "arb_Arab",
    "Chinese (Simplified)": "zho_Hans",
    "Japanese": "jpn_Jpan",
}

# ------------------------------------------------------------------ page setup
st.set_page_config(page_title="AIT Global Technologies", page_icon="🌍", layout="centered")

st.markdown("""
<style>
.stApp, [data-testid="stHeader"] { background: #FFFFFF; }
.stApp, .stApp p, .stApp label, .stApp span, .stApp li { color: #1B1B1B; }
.block-container { padding-top: 4.5rem !important; max-width: 820px; }

.title-card {
    background: #0B2A5B;
    border-radius: 12px;
    padding: 14px 22px;
    margin-bottom: 1.2rem;
    text-align: center;
}
.title-card h1 {
    margin: 0;
    padding: 0;
    font-size: 1.35rem;
    font-weight: 700;
    letter-spacing: 0.5px;
    white-space: nowrap;
    overflow: hidden;
    text-overflow: ellipsis;
}
.title-card h1 .ait  { color: #FFC107 !important; }
.title-card h1 .rest { color: #FFFFFF !important; }

.stButton > button, .stDownloadButton > button,
[data-testid="stFileUploader"] button {
    background: #FFC107 !important;
    color: #1B1B1B !important;
    border: none !important;
    border-radius: 8px !important;
    font-weight: 600 !important;
}
.stButton > button:hover, .stDownloadButton > button:hover {
    background: #FFB300 !important;
    color: #000000 !important;
}
.stDownloadButton > button { width: 100%; }
</style>

<div class="title-card">
    <h1>
        <span class="ait">AIT</span>
        <span class="rest"> GLOBAL TECHNOLOGIES</span>
    </h1>
</div>
""", unsafe_allow_html=True)

# ------------------------------------------------------------------ thread
def new_thread_id():
    return "THR-" + uuid.uuid4().hex[:8].upper()

if "thread_id" not in st.session_state:
    st.session_state.thread_id = new_thread_id()
if "streamed_text" not in st.session_state:
    st.session_state.streamed_text = ""
if "result" not in st.session_state:
    st.session_state.result = None

# ------------------------------------------------------------------ chromadb / history
@st.cache_resource
def get_collection():
    client = chromadb.PersistentClient(path=CHROMA_PATH, settings=Settings(anonymized_telemetry=False))
    return client.get_or_create_collection("translations")

def save_record(thread_id, source_name, src_lang, tgt_lang, source_text, translated_text):
    meta = {
        "thread_id": thread_id,
        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "source_name": source_name,
        "source_lang": src_lang,
        "target_lang": tgt_lang,
        "source_preview": source_text[:500],
        "chars": len(translated_text),
    }
    doc = translated_text[:5000]  # keep history light; full text is in the downloads
    if HAS_CHROMA:
        get_collection().add(ids=[uuid.uuid4().hex], documents=[doc], metadatas=[meta])
    else:
        st.session_state.setdefault("history", []).append(
            {"id": uuid.uuid4().hex, "doc": doc, "meta": meta}
        )

def parse_date_query(text):
    """Turn 'October 2026', 'oct 2026', '10/2026', '2026-10', '13 Oct 2026',
    '13/10/2026', '2026-10-13' or '2026' into a timestamp prefix like
    '2026-10' or '2026-10-13'. Returns None if the text is not a date."""
    s = re.sub(r"\s+", " ", text.replace(",", " ")).strip()
    if not s:
        return None

    if re.fullmatch(r"\d{4}", s):
        return s

    day_formats = ["%d %B %Y", "%d %b %Y", "%B %d %Y", "%b %d %Y",
                   "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"]
    month_formats = ["%B %Y", "%b %Y", "%m/%Y", "%m-%Y", "%Y-%m"]

    for fmt in day_formats:
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    for fmt in month_formats:
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m")
        except ValueError:
            pass
    return None

def load_records(thread_id=None, search=""):
    date_prefix = parse_date_query(search)

    if HAS_CHROMA:
        col = get_collection()
        where = {"thread_id": thread_id} if thread_id else None
        if search.strip() and not date_prefix:
            total = col.count()
            if total == 0:
                return []
            res = col.query(query_texts=[search.strip()], n_results=min(20, total), where=where)
            ids, docs, metas = res["ids"][0], res["documents"][0], res["metadatas"][0]
        else:
            res = col.get(where=where)
            ids, docs, metas = res["ids"], res["documents"], res["metadatas"]
        records = [{"id": i, "doc": d, "meta": m} for i, d, m in zip(ids, docs, metas)]
    else:
        records = st.session_state.get("history", [])
        if thread_id:
            records = [r for r in records if r["meta"]["thread_id"] == thread_id]
        if search.strip() and not date_prefix:
            q = search.strip().lower()
            records = [r for r in records if q in r["doc"].lower() or q in r["meta"]["source_preview"].lower()]

    if date_prefix:
        records = [r for r in records if r["meta"].get("timestamp", "").startswith(date_prefix)]

    records.sort(key=lambda r: r["meta"].get("timestamp", ""), reverse=True)
    return records[:50]

# ------------------------------------------------------------------ duplicate / noise cleaning
def collapse_loops(text, max_n=10):
    """Remove meaningless repetition loops the model sometimes produces
    (e.g. 'the house the house the house' or 'word word word word')."""
    out = []
    for w in text.split(" "):
        out.append(w)
        if len(out) >= 3 and out[-1] == out[-2] == out[-3]:
            del out[-1]                      # same word 3+ times in a row
            continue
        for n in range(2, max_n + 1):        # a phrase immediately repeated
            if len(out) >= 2 * n and out[-n:] == out[-2 * n:-n]:
                del out[-n:]
                break
    return " ".join(out)

def dedupe_lines(lines, remove_global):
    """Drop blank runs, lines with no letters (page numbers, '----'), consecutive
    duplicate lines, and (optionally) any line that already appeared earlier."""
    seen, out, prev = set(), [], None
    for line in lines:
        s = line.strip()
        if not s:
            if out and out[-1] != "":
                out.append("")
            continue
        if not re.search(r"[^\W\d_]", s):    # no letters at all -> no meaning
            continue
        key = re.sub(r"\s+", " ", s).lower()
        if key == prev:
            continue
        if remove_global and key in seen:
            continue
        seen.add(key)
        prev = key
        out.append(s)
    while out and out[-1] == "":
        out.pop()
    return out

# ------------------------------------------------------------------ model
@st.cache_resource(show_spinner="Loading model (first run only)...")
def load_model():
    path = snapshot_download(MODEL_NAME)
    tokenizer = AutoTokenizer.from_pretrained(path)
    translator = ctranslate2.Translator(path, device="cpu", compute_type="int8", inter_threads=1, intra_threads=2)
    return tokenizer, translator

def translate_batch(texts, src_code, tgt_code, clean=True):
    tokenizer, translator = load_model()
    tokenizer.src_lang = src_code
    batch = [tokenizer.convert_ids_to_tokens(tokenizer(t, truncation=True, max_length=256).input_ids) for t in texts]
    results = translator.translate_batch(
        batch,
        target_prefix=[[tgt_code]] * len(batch),
        beam_size=2,
        max_decoding_length=300,
        no_repeat_ngram_size=3,
        repetition_penalty=1.1,
    )
    outputs = []
    for r in results:
        tokens = r.hypotheses[0]
        if tokens and tokens[0] == tgt_code:
            tokens = tokens[1:]
        text = tokenizer.decode(tokenizer.convert_tokens_to_ids(tokens), skip_special_tokens=True)
        outputs.append(collapse_loops(text) if clean else text)
    return outputs

def translate_stream(text, src_code, tgt_code, on_partial=None, clean=True):
    """Translate one string, calling on_partial(partial_text) as tokens are generated.
    Falls back to normal batch translation if ctranslate2 has no `callback` support."""
    tokenizer, translator = load_model()
    tokenizer.src_lang = src_code
    source = tokenizer.convert_ids_to_tokens(tokenizer(text, truncation=True, max_length=256).input_ids)

    generated = []

    def decode(tokens):
        toks = [t for t in tokens if t != tgt_code]
        out = tokenizer.decode(tokenizer.convert_tokens_to_ids(toks), skip_special_tokens=True)
        return collapse_loops(out) if clean else out

    # Called from CTranslate2's worker thread: must NOT touch st.* elements.
    def callback(step):
        if step.token != tgt_code:
            generated.append(step.token)
        return None

    try:
        futures = translator.translate_batch(
            [source],
            target_prefix=[[tgt_code]],
            beam_size=1,
            max_decoding_length=300,
            no_repeat_ngram_size=3,
            repetition_penalty=1.1,
            callback=callback,
            asynchronous=True,
        )
    except TypeError:
        out = translate_batch([text], src_code, tgt_code, clean)[0]
        if on_partial:
            on_partial(out)
        return out

    future = futures[0]
    last_n = 0
    while not future.done():
        n = len(generated)
        if on_partial and n != last_n:
            on_partial(decode(list(generated)))
            last_n = n
        time.sleep(0.05)

    result = future.result()
    return decode(result.hypotheses[0])

# ------------------------------------------------------------------ helpers
def split_long(line):
    if len(line) <= MAX_CHUNK_CHARS:
        return [line]
    sentences = re.split(r"(?<=[.!?।።؟])\s+", line)
    chunks, cur = [], ""
    for s in sentences:
        if cur and len(cur) + len(s) + 1 > MAX_CHUNK_CHARS:
            chunks.append(cur)
            cur = s
        else:
            cur = f"{cur} {s}".strip()
    if cur:
        chunks.append(cur)
    return chunks

def read_upload(uploaded):
    name = uploaded.name.lower()
    data = uploaded.getvalue()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "txt", "\n\n".join((p.extract_text() or "") for p in reader.pages)
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        return "txt", "\n".join(p.text for p in doc.paragraphs)
    if name.endswith(".json"):
        return "json", json.loads(data.decode("utf-8", errors="replace"))
    return "txt", data.decode("utf-8", errors="replace")

def map_strings(obj, fn):
    if isinstance(obj, str):
        return fn(obj)
    if isinstance(obj, list):
        return [map_strings(x, fn) for x in obj]
    if isinstance(obj, dict):
        return {k: map_strings(v, fn) for k, v in obj.items()}
    return obj

def collect_strings(obj, acc):
    map_strings(obj, lambda s: acc.append(s) or s)
    return acc

def preview(text):
    if len(text) <= DISPLAY_LIMIT:
        return text
    return text[:DISPLAY_LIMIT] + "\n\n... (preview only - use the download buttons for the full text)"

# ------------------------------------------------------------------ PDF (self-contained Unicode support)
FONT_CACHE_DIR = os.path.join(tempfile.gettempdir(), "noto_fonts_cache")
NOTO_BASE = "https://github.com/notofonts/notofonts.github.io/raw/main/fonts/{n}/hinted/ttf/{n}-Regular.ttf"

BASE_FONT = ("NotoSans", NOTO_BASE.format(n="NotoSans"))

SCRIPT_FONTS = {
    "Hindi":     ("NotoDeva",  NOTO_BASE.format(n="NotoSansDevanagari")),
    "Marathi":   ("NotoDeva",  NOTO_BASE.format(n="NotoSansDevanagari")),
    "Telugu":    ("NotoTelu",  NOTO_BASE.format(n="NotoSansTelugu")),
    "Tamil":     ("NotoTaml",  NOTO_BASE.format(n="NotoSansTamil")),
    "Kannada":   ("NotoKnda",  NOTO_BASE.format(n="NotoSansKannada")),
    "Malayalam": ("NotoMlym",  NOTO_BASE.format(n="NotoSansMalayalam")),
    "Bengali":   ("NotoBeng",  NOTO_BASE.format(n="NotoSansBengali")),
    "Urdu":      ("NotoArab",  NOTO_BASE.format(n="NotoSansArabic")),
    "Arabic":    ("NotoArab",  NOTO_BASE.format(n="NotoSansArabic")),
    "Chinese (Simplified)": ("NotoSC", "https://github.com/google/fonts/raw/main/ofl/notosanssc/NotoSansSC%5Bwght%5D.ttf"),
    "Japanese":             ("NotoJP", "https://github.com/google/fonts/raw/main/ofl/notosansjp/NotoSansJP%5Bwght%5D.ttf"),
}
RTL_LANGS = {"Urdu", "Arabic"}


@st.cache_resource(show_spinner="Downloading PDF font (first time only)...")
def get_font_file(family, url):
    try:
        os.makedirs(FONT_CACHE_DIR, exist_ok=True)
        path = os.path.join(FONT_CACHE_DIR, f"{family}.ttf")
        if not (os.path.exists(path) and os.path.getsize(path) > 10_000):
            req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req, timeout=60) as resp, open(path, "wb") as f:
                f.write(resp.read())
        return path
    except Exception:
        return None


def _latin1_safe(text):
    return text.encode("latin-1", "replace").decode("latin-1")


@st.cache_data(show_spinner=False)
def make_pdf(lines, target_lang="English"):
    lines = list(lines)
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    unicode_ok = False
    base_family, base_url = BASE_FONT
    base_path = get_font_file(base_family, base_url)
    if base_path:
        try:
            pdf.add_font(base_family, "", base_path)
            pdf.set_font(base_family, size=11)
            unicode_ok = True
        except Exception:
            unicode_ok = False

    if unicode_ok:
        fallbacks = []
        if target_lang in SCRIPT_FONTS:
            fam, url = SCRIPT_FONTS[target_lang]
            p = get_font_file(fam, url)
            if p:
                try:
                    pdf.add_font(fam, "", p)
                    fallbacks.append(fam)
                except Exception:
                    pass
        if fallbacks:
            pdf.set_fallback_fonts(fallbacks)
        try:
            pdf.set_text_shaping(True)
        except Exception:
            pass
    else:
        pdf.set_font("Helvetica", size=11)

    align = "R" if target_lang in RTL_LANGS else "L"

    for line in lines:
        if not line.strip():
            pdf.ln(5)
            continue
        text = line if unicode_ok else _latin1_safe(line)
        try:
            pdf.multi_cell(pdf.epw, 6, text, new_x="LMARGIN", new_y="NEXT", align=align)
        except Exception:
            pdf.multi_cell(pdf.epw, 6, _latin1_safe(text), new_x="LMARGIN", new_y="NEXT")

    return bytes(pdf.output())


def make_docx(lines):
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

# ------------------------------------------------------------------ UI
tcol1, tcol2 = st.columns([3, 1])
tcol1.text_input("Thread ID (system-generated)", value=st.session_state.thread_id, disabled=True)
tcol2.markdown("<div style='height:1.75rem'></div>", unsafe_allow_html=True)
if tcol2.button("New thread"):
    st.session_state.thread_id = new_thread_id()
    st.session_state.result = None
    st.session_state.streamed_text = ""
    st.rerun()

col1, col2 = st.columns(2)
src_name = col1.selectbox("Translate from", list(LANGUAGES), index=0)
tgt_name = col2.selectbox("Translate to", list(LANGUAGES), index=2)  # Telugu

uploaded = st.file_uploader(f"Upload a file (max {MAX_UPLOAD_MB} MB)", type=["pdf", "txt", "docx", "json"])
if uploaded and uploaded.size > MAX_UPLOAD_MB * 1024 * 1024:
    st.error(f"File is {uploaded.size / 1024 / 1024:.1f} MB. The limit is {MAX_UPLOAD_MB} MB.")
    uploaded = None
typed = st.text_area("...or paste text", height=120)

oc1, oc2 = st.columns(2)
auto_clean = oc1.checkbox("Auto-remove meaningless repeats", value=True,
                          help="Removes repeated lines, number-only lines (page numbers) and repeated-word loops.")
remove_global = oc2.checkbox("Remove repeated lines anywhere in the document", value=True,
                             help="Untick if your document legitimately repeats the same line.")

st.subheader("Streaming Output")
output_placeholder = st.empty()
output_placeholder.text_area(
    "Output",
    preview(st.session_state.streamed_text),
    height=280,
    key="stream_display",
)

def show(text):
    """Push text into the Streaming Output box (fresh key each call so it always refreshes)."""
    st.session_state.streamed_text = text
    output_placeholder.text_area("Output", preview(text), height=280, key=f"out_{uuid.uuid4().hex}")

if st.button("Translate", type="primary"):
    if src_name == tgt_name:
        st.info("Source and target languages are the same.")
    elif not uploaded and not typed.strip():
        st.warning("Upload a file or paste some text first.")
    else:
        st.session_state.result = None
        st.session_state.streamed_text = ""

        try:
            kind, content = read_upload(uploaded) if uploaded else ("txt", typed)
        except Exception as e:
            st.error(f"Could not read the file: {e}")
            st.stop()

        src, tgt = LANGUAGES[src_name], LANGUAGES[tgt_name]
        bar = st.progress(0.0, text="Translating...")

        if kind == "json":
            strings = collect_strings(content, [])
            unique = list(dict.fromkeys(s for s in strings if s.strip()))
            done = {}
            accumulated = []

            for n, s in enumerate(unique):
                base = "\n".join(accumulated[-200:])
                t = translate_stream(
                    s, src, tgt,
                    on_partial=lambda p, base=base: show(f"{base}\n{p}" if base else p),
                    clean=auto_clean,
                )
                done[s] = t
                accumulated.append(t)
                show("\n".join(accumulated[-200:]))
                bar.progress((n + 1) / max(len(unique), 1))

            lines = [done.get(s, s) for s in strings]
            source_text = "\n".join(strings)
            result_json = map_strings(content, lambda s: done.get(s, s))
        else:
            src_lines = content.splitlines()
            if auto_clean:
                src_lines = dedupe_lines(src_lines, remove_global)
                content = "\n".join(src_lines)

            pieces = []
            for l in src_lines:
                chunks = split_long(l.strip()) if l.strip() else []
                if auto_clean:  # drop sentences repeated back-to-back inside one line
                    chunks = [c for i, c in enumerate(chunks) if i == 0 or c != chunks[i - 1]]
                pieces.append(chunks)

            flat = [p for chunks in pieces for p in chunks]
            unique = list(dict.fromkeys(s for s in flat if s.strip()))
            done = {}

            def render(extra=None):
                m = {**done, **(extra or {})}
                rows = [" ".join(m[p] for p in chunks if p in m) for chunks in pieces]
                return "\n".join(rows).rstrip()

            if len(unique) > LARGE_THRESHOLD:
                # Large file: batched translation (shortest first to reduce padding)
                order = sorted(unique, key=len)
                for i in range(0, len(order), BIG_BATCH):
                    batch = order[i:i + BIG_BATCH]
                    outs = translate_batch(batch, src, tgt, clean=auto_clean)
                    done.update(zip(batch, outs))
                    if (i // BIG_BATCH) % 10 == 0:
                        show("(Large file - showing the latest part only)\n..." + render()[-3000:])
                    bar.progress(min((i + BIG_BATCH) / len(order), 1.0))
            else:
                for n, s in enumerate(unique):
                    t = translate_stream(
                        s, src, tgt,
                        on_partial=lambda p, s=s: show(render({s: p})),
                        clean=auto_clean,
                    )
                    done[s] = t
                    show(render())
                    bar.progress((n + 1) / max(len(unique), 1))

            lines = [" ".join(done.get(p, p) for p in chunks) if chunks else "" for chunks in pieces]
            source_text = content
            result_json = {
                "thread_id": st.session_state.thread_id,
                "source_language": src_name,
                "target_language": tgt_name,
                "translated_text": "\n".join(lines),
            }

        bar.empty()
        final_text = "\n".join(lines)
        show(final_text)
        st.session_state.result = {"lines": lines, "json": result_json, "target": tgt_name}

        try:
            save_record(
                st.session_state.thread_id,
                uploaded.name if uploaded else "pasted text",
                src_name, tgt_name, source_text, final_text,
            )
        except Exception as e:
            st.warning(f"Translation done, but history could not be saved: {e}")

# ------------------------------------------------------------------ downloads
result = st.session_state.result
if result:
    st.markdown("---")
    st.subheader("Download translated file")
    c1, c2, c3, c4 = st.columns(4)

    try:
        pdf_bytes = make_pdf(tuple(result["lines"]), result.get("target", "English"))
        c1.download_button("📄 PDF", pdf_bytes, "translation.pdf", "application/pdf", use_container_width=True)
    except Exception as e:
        c1.button("📄 PDF (error)", disabled=True, use_container_width=True)
        st.warning(f"PDF could not be generated: {e}")

    c2.download_button("📝 TXT", "\n".join(result["lines"]).encode("utf-8"), "translation.txt", "text/plain", use_container_width=True)
    c3.download_button("📑 DOCX", make_docx(result["lines"]), "translation.docx",
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
    c4.download_button("📋 JSON", json.dumps(result["json"], ensure_ascii=False, indent=2).encode("utf-8"),
                       "translation.json", "application/json", use_container_width=True)

# ------------------------------------------------------------------ history
st.divider()
st.subheader("Past records")
if not HAS_CHROMA:
    st.caption("History is session-only (ChromaDB not installed).")

h1, h2 = st.columns([1, 2])
scope = h1.radio("Show", ["This thread", "All threads"], horizontal=True)
search = h2.text_input("Search", placeholder="Word, or a date like 'October 2026' / '13 Oct 2026'")

try:
    records = load_records(st.session_state.thread_id if scope == "This thread" else None, search)
except Exception as e:
    records = []
    st.warning(f"Could not load records: {e}")

if not records:
    st.caption("No saved records yet.")
else:
    for r in records:
        m = r["meta"]
        title = f"{m.get('timestamp', '')}  ·  {m.get('source_name', '')}  ·  {m.get('source_lang', '')} → {m.get('target_lang', '')}"
        with st.expander(title):
            st.markdown("**Source (preview)**")
            st.text(m.get("source_preview", ""))
            st.markdown("**Translation**")
            st.text(r["doc"])